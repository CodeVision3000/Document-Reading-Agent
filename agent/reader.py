"""
Document reader module - handles reading various document formats
"""
import os


class DocumentReader:
    """Reads documents of various formats and returns their text content."""

    SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx", ".xlsx"}

    def read(self, file_path: str) -> str:
        """
        Read a document and return its text content.

        Args:
            file_path: Path to the document file.

        Returns:
            The text content of the document.

        Raises:
            FileNotFoundError: If the file does not exist.
            ValueError: If the file format is not supported.
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        _, ext = os.path.splitext(file_path)
        ext = ext.lower()

        if ext not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported file format: {ext}. Supported formats: {self.SUPPORTED_EXTENSIONS}")

        if ext == ".txt":
            return self._read_txt(file_path)
        elif ext == ".pdf":
            return self._read_pdf(file_path)
        elif ext == ".docx":
            return self._read_docx(file_path)
        elif ext == ".xlsx":
            return self._read_xlsx(file_path)

        raise ValueError(f"Unsupported file format: {ext}")

    def _read_txt(self, file_path: str) -> str:
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()

    def _read_pdf(self, file_path: str) -> str:
        try:
            from pypdf import PdfReader
            reader = PdfReader(file_path)
            return "\n".join(page.extract_text() for page in reader.pages)
        except ImportError:
            raise ImportError("pypdf is required to read PDF files. Install it with: pip install pypdf")

    def _read_docx(self, file_path: str) -> str:
        try:
            import docx
            doc = docx.Document(file_path)
            return "\n".join(paragraph.text for paragraph in doc.paragraphs)
        except ImportError:
            raise ImportError("python-docx is required to read DOCX files. Install it with: pip install python-docx")

    def _read_xlsx(self, file_path: str) -> str:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            lines = []
            for sheet in wb.worksheets:
                lines.append(f"[Sheet: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    row_text = " | ".join("" if cell is None else str(cell).replace(" | ", " ") for cell in row)
                    if row_text.strip():
                        lines.append(row_text)
            return "\n".join(lines)
        except ImportError:
            raise ImportError("openpyxl is required to read XLSX files. Install it with: pip install openpyxl")
